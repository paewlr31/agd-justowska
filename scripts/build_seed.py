# -*- coding: utf-8 -*-
"""Resize the owner's photos and build data/seed.json from the Excel catalog."""
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent
EXCEL = SOURCE / "producenci" / "Modele do strony internetowej JUSTOWSKA.xlsx"
PUBLIC = ROOT / "public"
STAMP = "2026-10-01T08:00:00.000Z"

BRANDS = [
    "AEG",
    "ASKO",
    "BORA",
    "CIARKO DESIGN",
    "COSENTINO",
    "ELECTROLUX",
    "ELICA",
    "FRANKE",
    "FALMEC",
    "FLORIM",
    "KERNAU",
    "MIELE",
    "LIEBHERR",
    "NORDBERG",
    "REGINOX",
    "SMEG",
    "SCHOCK",
    "SIEMENS",
    "QUOOKER",
]

USER_CATEGORIES = {
    "AEG": ["Piekarnik", "Mikrofala", "Płyta indukcyjna", "Okap", "Ekspres ciśnieniowy", "Lodówka", "Winiarka", "Zmywarka", "Pralka", "Suszarka bębnowa", "Pralko-suszarka"],
    "ASKO": ["Piekarnik", "Piekarnik z mikrofalą", "Mikrofala", "Płyta gazowa", "Płyta indukcyjna", "Zmywarka", "Lodówka", "Pralka", "Suszarka bębnowa", "Ekspres ciśnieniowy"],
    "CIARKO DESIGN": ["Okap przyścienny", "Okap podszafkowy", "Okap wyspowy", "Okap blatowy"],
    "ELECTROLUX": ["Płyta indukcyjna z wyciągiem", "Płyta indukcyjna", "Płyta gazowa", "Piekarnik", "Mikrofala", "Ekspres ciśnieniowy", "Lodówka", "Zamrażarka", "Okap", "Zmywarka", "Pralka", "Suszarka", "Pralko-suszarka"],
    "LIEBHERR": ["Lodówka do zabudowy", "Chłodziarka do zabudowy", "Zamrażarka do zabudowy"],
    "SMEG": ["Piekarnik", "Mikrofala", "Płyta gazowa", "Płyta indukcyjna", "Ekspres ciśnieniowy", "Zmywarka do zabudowy", "Blast chiller", "Lodówka wolnostojąca"],
    "SCHOCK": ["Zlewozmywak granitowy", "Bateria z wyciąganą wylewką", "Bateria bez wyciąganej wylewki"],
}

CANON = {
    "piekarnik": "Piekarnik",
    "mikrofala": "Mikrofala",
    "plyta indukcyjna": "Płyta indukcyjna",
    "okap": "Okap",
    "ekspres cisnieniowy": "Ekspres ciśnieniowy",
    "lodowka": "Lodówka",
    "winiarka": "Winiarka",
    "zmywarka": "Zmywarka",
    "pralka": "Pralka",
    "suszarka bebnowa": "Suszarka bębnowa",
    "pralko-suszarka": "Pralko-suszarka",
    "pralkosuszarka": "Pralko-suszarka",
    "piekarnik z mikrofala": "Piekarnik z mikrofalą",
    "plyta gazowa": "Płyta gazowa",
    "okap przyscienny": "Okap przyścienny",
    "okap podszafkowy": "Okap podszafkowy",
    "okap wyspowy": "Okap wyspowy",
    "okap blatowy": "Okap blatowy",
    "plyta indukcyjna z wyciagiem": "Płyta indukcyjna z wyciągiem",
    "zamrazarka": "Zamrażarka",
    "suszarka": "Suszarka",
    "lodowka do zabudowy": "Lodówka do zabudowy",
    "chlodziarka do zabudowy": "Chłodziarka do zabudowy",
    "zamrazarka do zabudowy": "Zamrażarka do zabudowy",
    "zlewozmywak granitowy": "Zlewozmywak granitowy",
    "bateria z wyciagana wylewka": "Bateria z wyciąganą wylewką",
    "bateria bez wyciaganej wylewki": "Bateria bez wyciąganej wylewki",
    "zmywarka do zabudowy": "Zmywarka do zabudowy",
    "blast chiller": "Blast chiller",
    "lodowka wolnostojaca": "Lodówka wolnostojąca",
}


def fold(value: str) -> str:
    text = unicodedata.normalize("NFKD", value.strip().lower())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("ł", "l")
    text = re.sub(r"\s+", " ", text)
    return text


def canon(value: str) -> str:
    key = fold(value)
    if key not in CANON:
        raise SystemExit(f"Nieznana kategoria: {value!r} -> {key!r}")
    return CANON[key]


def slugify(value: str) -> str:
    text = fold(value)
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text


def save_jpeg(src: Path, dest: Path, max_width: int, quality: int) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(src) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        if image.width > max_width:
            height = round(image.height * (max_width / image.width))
            image = image.resize((max_width, height), Image.Resampling.LANCZOS)
        image.save(dest, "JPEG", quality=quality, optimize=True)


def save_logo() -> None:
    with Image.open(SOURCE / "logo" / "logo.png") as image:
        image = image.convert("RGBA")
        cropped = image.crop(image.getbbox())
        cropped.thumbnail((640, 640), Image.Resampling.LANCZOS)
        dest = PUBLIC / "logo.png"
        dest.parent.mkdir(parents=True, exist_ok=True)
        cropped.save(dest, "PNG", optimize=True)

        icon = Image.new("RGBA", (256, 256), (128, 0, 32, 255))
        mark = cropped.copy()
        mark.thumbnail((210, 210), Image.Resampling.LANCZOS)
        icon.paste(mark, ((256 - mark.width) // 2, (256 - mark.height) // 2), mark)
        icon_dest = ROOT / "app" / "icon.png"
        icon.convert("RGB").save(icon_dest, "PNG", optimize=True)


def load_products() -> list[dict]:
    workbook = openpyxl.load_workbook(EXCEL, data_only=True)
    products = []
    seen = set()
    for sheet in workbook.sheetnames:
        for row in workbook[sheet].iter_rows(values_only=True):
            brand, model, category, price, description = (list(row) + [None] * 5)[:5]
            if not model or str(brand).strip() in {"Marka", "Opisy do strony WWW"}:
                continue
            brand_name = str(brand).strip()
            model_name = str(model).strip()
            category_name = canon(str(category))
            if isinstance(price, (int, float)):
                price_value = int(price) if float(price).is_integer() else round(float(price), 2)
            elif price in (None, ""):
                price_value = None
            else:
                raise SystemExit(f"Zła cena: {brand_name} {model_name} {price!r}")
            text = re.sub(r"\s+", " ", str(description or "")).strip()
            text = re.sub(r"\s*-\s*Autoryzowany sklep Ciarko Design\s*$", "", text, flags=re.I).strip()
            if text.lower().startswith("http://") or text.lower().startswith("https://"):
                text = ""
            product_id = f"{slugify(brand_name)}-{slugify(model_name)}"
            if product_id in seen:
                raise SystemExit(f"Duplikat id: {product_id}")
            seen.add(product_id)
            products.append({
                "id": product_id,
                "brand": brand_name,
                "model": model_name,
                "category": category_name,
                "price": price_value,
                "description": text,
                "image": None,
                "createdAt": STAMP,
            })
    return products


def main() -> None:
    save_logo()
    home_src = SOURCE / "homePagePhotos"
    home_map = {
        "Obraz1.jpg": ("ovens.jpg", 1400),
        "Obraz2.jpg": ("hero.jpg", 2000),
        "Obraz3.jpg": ("hob.jpg", 1400),
        "Obraz4.jpg": ("tap.jpg", 1400),
        "Obraz5.jpg": ("sink.jpg", 1200),
    }
    for name, (dest, width) in home_map.items():
        save_jpeg(home_src / name, PUBLIC / "media" / "home" / dest, width, 84)

    gallery = []
    files = sorted((SOURCE / "galeria").glob("*.jpg"))
    for index, src in enumerate(files, start=1):
        filename = f"{index:02d}.jpg"
        save_jpeg(src, PUBLIC / "media" / "galeria" / filename, 1600, 80)
        gallery.append({
            "id": f"galeria-{index:02d}",
            "image": f"/media/galeria/{filename}",
            "caption": "",
            "createdAt": STAMP,
        })

    products = load_products()
    by_brand: dict[str, list[str]] = {brand: [] for brand in BRANDS}
    for product in products:
        if product["brand"] not in by_brand:
            raise SystemExit(f"Marka spoza listy: {product['brand']}")
        if product["category"] not in by_brand[product["brand"]]:
            by_brand[product["brand"]].append(product["category"])

    manufacturers = []
    for order, brand in enumerate(BRANDS):
        categories = []
        for name in USER_CATEGORIES.get(brand, []):
            if name not in categories:
                categories.append(name)
        for name in by_brand[brand]:
            if name not in categories:
                categories.append(name)
        manufacturers.append({
            "slug": slugify(brand),
            "name": brand,
            "categories": categories,
            "sortOrder": order,
        })

    seed = {"manufacturers": manufacturers, "products": products, "gallery": gallery}
    dest = ROOT / "data" / "seed.json"
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(seed, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"products {len(products)} gallery {len(gallery)} brands {len(manufacturers)}")
    print("generated", datetime.now(timezone.utc).isoformat())


if __name__ == "__main__":
    main()
